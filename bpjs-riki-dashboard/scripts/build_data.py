#!/usr/bin/env python3
"""
Parse "LAPORAN HARIAN RIKI" (PIC BPJS) workbook into src/data/dashboard-data.json.

Usage:
    python3 scripts/build_data.py <path-ke-xlsx> [output.json]

Setiap sheet bulan (APRIL, MEI, JUNI, JULI, AGUSTUS, ...) berisi dua bagian:
  1. Tabel "PROGRESS KERJA" -- satu baris per lokasi/klien, checklist per tahap
     BPJS Kesehatan (KES) & Ketenagakerjaan (TK), sampai baris "TOTAL".
  2. Bagian "LAPORAN HARIAN" -- log aktivitas harian (TGL, KEGIATAN, KETERANGAN).

Kolom tabel progress (tetap sama di semua sheet bulan, kolom 1-18):
  1 NO | 2 kelompok KES (jarang diisi, berlaku turun sampai nilai baru) |
  3 LOKASI (nama entitas, selalu diisi) |
  4 rekapGaji.kes | 5 rekapGaji.tk |
  6 rekapBayar.kes | 7 rekapBayar.tk |
  8 (spacer) | 9 nominalBayar.kes | 10 (spacer) | 11 nominalBayar.tk |
  12 pengecekanGaji.tk | 13 finalisasi.tk |
  14 buktiBayar.kes | 15 buktiBayar.tk |
  16 keterangan.kes | 17 keterangan.tk | 18 note
"""
import json
import sys
from datetime import datetime, date

import openpyxl

MONTH_SHEETS = [
    ("APRIL", 4),
    ("MEI", 5),
    ("JUNI", 6),
    ("JULI", 7),
    ("AGUSTUS", 8),
]
YEAR = 2026

STAGES = [
    ("rekapGaji", "REKAP U/ PENGGAJIAN", True),
    ("rekapBayar", "REKAP PEMBAYARAN", True),
    ("nominalBayar", "NOMINAL BAYAR", True),
    ("pengecekanGaji", "PENGECEKAN REKAP GAJI", False),
    ("finalisasi", "FINALISASI", False),
    ("buktiBayar", "BUKTI BAYAR BPJS", True),
]
# (key, label, hasKesColumn) -- pengecekanGaji & finalisasi hanya melacak TK


def clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def parse_progress_sheet(ws, month_code, month_label):
    locations = []
    kelompok_kes = None
    r = 5
    while r <= ws.max_row:
        a = ws.cell(row=r, column=1).value
        b = clean(ws.cell(row=r, column=2).value)
        c = clean(ws.cell(row=r, column=3).value)
        if isinstance(c, str) and c.upper() == "TOTAL":
            r += 1
            break
        if isinstance(b, str) and b.upper() == "LAPORAN HARIAN":
            break
        if a is None and b is None and c is None:
            r += 1
            continue
        if b:
            kelompok_kes = b
        if c:
            def bval(col):
                v = ws.cell(row=r, column=col).value
                return v if isinstance(v, bool) else None

            stages = {}
            stages["rekapGaji"] = {"kes": bval(4), "tk": bval(5)}
            stages["rekapBayar"] = {"kes": bval(6), "tk": bval(7)}
            stages["nominalBayar"] = {"kes": bval(9), "tk": bval(11)}
            stages["pengecekanGaji"] = {"tk": bval(12)}
            stages["finalisasi"] = {"tk": bval(13)}
            stages["buktiBayar"] = {"kes": bval(14), "tk": bval(15)}

            cells = []
            for key, _label, has_kes in STAGES:
                if has_kes:
                    cells.append(stages[key].get("kes"))
                cells.append(stages[key].get("tk"))
            applicable = [v for v in cells if v is not None]
            if not applicable:
                status = "no-data"
            elif all(applicable):
                status = "selesai"
            elif not any(applicable):
                status = "belum"
            else:
                status = "proses"
            progress_pct = (
                round(100 * sum(1 for v in applicable if v) / len(applicable), 1)
                if applicable
                else None
            )

            locations.append(
                {
                    "month": month_code,
                    "no": a,
                    "kelompokKes": kelompok_kes,
                    "lokasi": c,
                    "stages": stages,
                    "keteranganKes": clean(ws.cell(row=r, column=16).value),
                    "keteranganTk": clean(ws.cell(row=r, column=17).value),
                    "note": clean(ws.cell(row=r, column=18).value),
                    "status": status,
                    "progressPct": progress_pct,
                }
            )
        r += 1
    return locations, r


def parse_daily_log(ws, start_row, month_code):
    entries = []
    r = start_row
    # skip section title + header row ("LAPORAN HARIAN" / "TGL | KEGIATAN | KETERANGAN")
    while r <= ws.max_row:
        a = clean(ws.cell(row=r, column=1).value)
        if isinstance(a, str) and a.upper() == "TGL":
            r += 1
            break
        r += 1
    current_date = None
    while r <= ws.max_row:
        raw_date = ws.cell(row=r, column=1).value
        kegiatan = clean(ws.cell(row=r, column=2).value)
        keterangan = clean(ws.cell(row=r, column=3).value)
        parsed_date = parse_log_date(raw_date, fix_mmdd_swap=True)
        if parsed_date:
            current_date = parsed_date
        if kegiatan:
            entries.append(
                {
                    "date": current_date,
                    "kegiatan": kegiatan,
                    "keterangan": keterangan,
                    "sourceMonth": month_code,
                }
            )
        r += 1
    return entries


def parse_log_date(raw, fix_mmdd_swap=False):
    if isinstance(raw, (datetime, date)):
        # Entries in the per-month "LAPORAN HARIAN" sections are typed as
        # dd/mm text; whenever the day is <=12, Excel silently auto-converts
        # that text to a real date using MM/DD instead, swapping the two
        # fields (e.g. "12/08" meant 12 Agustus becomes 8 Desember). Since a
        # genuine month can never exceed 12, this swap is only ever possible
        # -- and only ever happened in this data -- when raw.day <= 12, so
        # undo it in that case. Sheet1's dates are entered directly (not
        # auto-converted from dd/mm text) and must NOT be swapped.
        if fix_mmdd_swap and raw.day <= 12:
            return date(raw.year, raw.day, raw.month).strftime("%Y-%m-%d")
        return raw.strftime("%Y-%m-%d")
    if isinstance(raw, str):
        raw = raw.strip()
        if not raw:
            return None
        for fmt in ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d/%m"):
            try:
                dt = datetime.strptime(raw, fmt)
                year = dt.year if dt.year > 1904 else YEAR
                return date(year, dt.month, dt.day).strftime("%Y-%m-%d")
            except ValueError:
                continue
    return None


def parse_payments(ws):
    payments = []
    for r in range(5, ws.max_row + 1):
        no = ws.cell(row=r, column=1).value
        lokasi = clean(ws.cell(row=r, column=2).value)
        kes = ws.cell(row=r, column=3).value
        tk = ws.cell(row=r, column=4).value
        if not lokasi:
            continue
        nominal_kes = kes if isinstance(kes, (int, float)) else None
        nominal_tk = tk if isinstance(tk, (int, float)) else None
        payments.append(
            {
                "no": no,
                "lokasi": lokasi,
                "nominalKes": nominal_kes,
                "nominalKesNote": kes if isinstance(kes, str) else None,
                "nominalTk": nominal_tk,
                "total": (nominal_kes or 0) + (nominal_tk or 0),
            }
        )
    return payments


def build_month_summary(month_code, month_label, locations):
    stage_counts = {}
    for key, _label, has_kes in STAGES:
        sub_keys = ["kes", "tk"] if has_kes else ["tk"]
        stage_counts[key] = {}
        for sub in sub_keys:
            selesai = sum(1 for loc in locations if loc["stages"][key].get(sub) is True)
            belum = sum(1 for loc in locations if loc["stages"][key].get(sub) is False)
            kosong = sum(1 for loc in locations if loc["stages"][key].get(sub) is None)
            stage_counts[key][sub] = {"selesai": selesai, "belum": belum, "kosong": kosong}

    pct_values = [loc["progressPct"] for loc in locations if loc["progressPct"] is not None]
    overall_pct = round(sum(pct_values) / len(pct_values), 1) if pct_values else None
    return {
        "code": month_code,
        "label": month_label,
        "totalLokasi": len(locations),
        "lokasiSelesai": sum(1 for loc in locations if loc["status"] == "selesai"),
        "lokasiProses": sum(1 for loc in locations if loc["status"] == "proses"),
        "lokasiBelum": sum(1 for loc in locations if loc["status"] == "belum"),
        "lokasiNoData": sum(1 for loc in locations if loc["status"] == "no-data"),
        "overallPct": overall_pct,
        "stages": stage_counts,
    }


def summarize_notes(locations):
    counter = {}
    for loc in locations:
        for text in (loc["note"], loc["keteranganKes"], loc["keteranganTk"]):
            if not text:
                continue
            key = text.strip()
            entry = counter.setdefault(key, {"text": key, "count": 0, "locations": []})
            entry["count"] += 1
            if loc["lokasi"] not in entry["locations"]:
                entry["locations"].append(loc["lokasi"])
    return sorted(counter.values(), key=lambda e: -e["count"])


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else "src/data/dashboard-data.json"

    wb = openpyxl.load_workbook(src_path, data_only=True)

    months = []
    all_locations = []
    all_daily_log = []

    for sheet_name, month_num in MONTH_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        month_code = f"{YEAR}-{month_num:02d}"
        month_label = f"{sheet_name.capitalize()} {YEAR}"
        locations, log_start_row = parse_progress_sheet(ws, month_code, month_label)
        daily_log = parse_daily_log(ws, log_start_row, month_code)
        months.append(build_month_summary(month_code, month_label, locations))
        all_locations.extend(locations)
        all_daily_log.extend(daily_log)

    # Fold in the standalone consolidated log (Sheet1) for entries not already
    # covered by a per-month "LAPORAN HARIAN" section (e.g. dates before April
    # sheet's own log starts, or duplicate-free extra detail).
    if "Sheet1" in wb.sheetnames:
        ws1 = wb["Sheet1"]
        current_date = None
        seen = {(e["date"], e["kegiatan"]) for e in all_daily_log}
        for r in range(4, ws1.max_row + 1):
            raw_date = ws1.cell(row=r, column=2).value
            kegiatan = clean(ws1.cell(row=r, column=3).value)
            keterangan = clean(ws1.cell(row=r, column=4).value)
            parsed = parse_log_date(raw_date, fix_mmdd_swap=False)
            if parsed:
                current_date = parsed
            if kegiatan and (current_date, kegiatan) not in seen:
                all_daily_log.append(
                    {
                        "date": current_date,
                        "kegiatan": kegiatan,
                        "keterangan": keterangan,
                        "sourceMonth": None,
                    }
                )
                seen.add((current_date, kegiatan))

    all_daily_log = [e for e in all_daily_log if e["date"]]
    all_daily_log.sort(key=lambda e: e["date"])

    payments = parse_payments(wb["Sheet3"]) if "Sheet3" in wb.sheetnames else []
    notes = summarize_notes(all_locations)

    dataset = {
        "generatedAt": datetime.now().strftime("%Y-%m-%d"),
        "months": months,
        "locations": all_locations,
        "dailyLog": all_daily_log,
        "payments": payments,
        "notes": notes,
    }

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(dataset, f, ensure_ascii=False, indent=2)

    print(f"OK -> {out_path}")
    print(f"  months: {[m['code'] for m in months]}")
    print(f"  total lokasi rows: {len(all_locations)}")
    print(f"  daily log entries: {len(all_daily_log)}")
    print(f"  payments rows: {len(payments)}")
    print(f"  distinct notes: {len(notes)}")


if __name__ == "__main__":
    main()
