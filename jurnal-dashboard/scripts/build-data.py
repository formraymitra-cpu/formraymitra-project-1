#!/usr/bin/env python3
"""
Mengubah "JURNAL HARIAN DINI SAFFANAH 2026.xlsx" (hasil gabungan, satu sheet
per bulan: TANGGAL|HARI|JAM MASUK|JAM PULANG|NO|DAILY WORK PLAN|CEKLIST|
TIME SCHEDULE|KETERANGAN, header di baris 3, data mulai baris 4) menjadi satu
dataset JSON yang dikonsumsi dashboard (src/data/jurnal-data.json).

Jumlah foto per tanggal dihitung otomatis dari header sheet galeri
"<BULAN> FOTO" ("DD/MM/YYYY (Hari) - N foto"), bukan dari kolom manual di
sheet data — jadi selalu sinkron dengan isi galeri yang sebenarnya.

Jalankan ulang setiap kali ada bulan baru digabungkan:
    python3 scripts/build-data.py <path-ke-jurnal.xlsx> [path-ke-invoice.xlsx]

Argumen kedua (opsional) adalah workbook "MONITORING INVOICE DINI" — kalau
diisi, hasilnya digabung sebagai key "invoice" di dataset yang sama.
"""
import json
import re
import sys
from datetime import datetime, date

import openpyxl

MONTH_LABEL_ID = {
    1: "Januari", 2: "Februari", 3: "Maret", 4: "April", 5: "Mei", 6: "Juni",
    7: "Juli", 8: "Agustus", 9: "September", 10: "Oktober", 11: "November", 12: "Desember",
}
MONTH_NAMES_UPPER = {v.upper() for v in MONTH_LABEL_ID.values()}
MONTH_NUM_ID = {v.upper(): k for k, v in MONTH_LABEL_ID.items()}
PLACEHOLDER_TASK = "(TIDAK ADA DATA)"


def iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None


def cell_display(v):
    """Sebagian sel kontrak/SP diisi tanggal Excel asli, sebagian lagi
    diketik manual sebagai teks ("29 Juni 2026") dan bisa multi-baris kalau
    satu lokasi punya beberapa SP/BAPP sekaligus. Fungsi ini menyeragamkan
    keduanya jadi satu string siap-tampil (bukan ISO date, karena sumbernya
    campur teks & tanggal)."""
    if v is None or v == "":
        return None
    if isinstance(v, (datetime, date)):
        return f"{v.day} {MONTH_LABEL_ID[v.month]} {v.year}"
    s = str(v).strip()
    return s if s else None


def parse_hhmm(s):
    if not s or not isinstance(s, str) or ":" not in s:
        return None
    try:
        h, m = s.split(":")
        return int(h) + int(m) / 60
    except ValueError:
        return None


def build_foto_count_map(wb):
    """Hitung jumlah foto per tanggal langsung dari header sheet galeri
    "<BULAN> FOTO" ("DD/MM/YYYY (Hari) - N foto" di kolom A) — bukan dari
    kolom JUMLAH FOTO manual di sheet data, supaya selalu sinkron dengan isi
    galeri yang sebenarnya."""
    count_by_date = {}
    pattern = re.compile(r"^(\d{1,2}/\d{1,2}/\d{4})\b.*?(\d+)\s*foto", re.IGNORECASE)
    for sn in wb.sheetnames:
        if not sn.strip().upper().endswith(" FOTO"):
            continue
        ws = wb[sn]
        for row in ws.iter_rows(min_col=1, max_col=1, values_only=True):
            v = row[0]
            if not isinstance(v, str):
                continue
            m = pattern.match(v)
            if m:
                d, mth, y = m.group(1).split("/")
                iso = f"{y}-{int(mth):02d}-{int(d):02d}"
                count_by_date[iso] = int(m.group(2))
    return count_by_date


def parse_month_sheet(ws, foto_count_map):
    tasks = []
    r = 4
    while True:
        tgl = ws.cell(r, 1).value
        if tgl is None:
            break
        if isinstance(tgl, datetime):
            tgl_date = tgl.date()
        elif isinstance(tgl, date):
            tgl_date = tgl
        else:
            r += 1
            continue
        tgl_iso = tgl_date.isoformat()
        hari = ws.cell(r, 2).value
        jam_masuk = ws.cell(r, 3).value
        jam_pulang = ws.cell(r, 4).value
        no = ws.cell(r, 5).value
        plan = ws.cell(r, 6).value
        ceklist = ws.cell(r, 7).value
        jadwal = ws.cell(r, 8).value
        keterangan = ws.cell(r, 9).value
        tasks.append({
            "tanggal": tgl_iso,
            "hari": hari,
            "jamMasuk": jam_masuk,
            "jamPulang": jam_pulang,
            "no": no,
            "tugas": plan,
            "selesai": ceklist is True,
            "jadwal": jadwal,
            "keterangan": keterangan,
            "jumlahFoto": foto_count_map.get(tgl_iso, 0),
        })
        r += 1
    return tasks


def build_days(all_tasks):
    by_date = {}
    for t in all_tasks:
        by_date.setdefault(t["tanggal"], []).append(t)

    days = []
    for tgl, items in by_date.items():
        real_items = [t for t in items if t["tugas"] and str(t["tugas"]).strip().upper() != PLACEHOLDER_TASK]
        total = len(real_items)
        selesai = sum(1 for t in real_items if t["selesai"])
        jam_masuk_h = parse_hhmm(items[0]["jamMasuk"])
        jam_pulang_h = parse_hhmm(items[0]["jamPulang"])
        durasi = None
        if jam_masuk_h is not None and jam_pulang_h is not None and jam_pulang_h >= jam_masuk_h:
            durasi = round(jam_pulang_h - jam_masuk_h, 2)
        days.append({
            "tanggal": tgl,
            "hari": items[0]["hari"],
            "jamMasuk": items[0]["jamMasuk"],
            "jamPulang": items[0]["jamPulang"],
            "jamKerjaJam": durasi,
            "totalTugas": total,
            "selesai": selesai,
            "pctSelesai": round(selesai / total, 4) if total else None,
            "jumlahFoto": items[0]["jumlahFoto"],
        })
    days.sort(key=lambda d: d["tanggal"])
    return days


def build_months(days):
    by_month = {}
    for d in days:
        y, m, _ = d["tanggal"].split("-")
        code = f"{y}-{m}"
        by_month.setdefault(code, []).append(d)

    months = []
    for code in sorted(by_month.keys()):
        items = by_month[code]
        y, m = code.split("-")
        total_tugas = sum(d["totalTugas"] for d in items)
        selesai = sum(d["selesai"] for d in items)
        jam_list = [d["jamKerjaJam"] for d in items if d["jamKerjaJam"] is not None]
        months.append({
            "code": code,
            "label": f"{MONTH_LABEL_ID[int(m)]} {y}",
            "hariTercatat": len(items),
            "totalTugas": total_tugas,
            "selesai": selesai,
            "pctSelesai": round(selesai / total_tugas, 4) if total_tugas else None,
            "rataJamKerja": round(sum(jam_list) / len(jam_list), 2) if jam_list else None,
            "totalFoto": sum(d["jumlahFoto"] for d in items),
        })
    return months


def parse_tagihan_sheet(ws):
    tagihan = []
    notes = []
    if ws is None:
        return tagihan, notes
    current = None
    blank_streak = 0
    r = 5
    last_row = ws.max_row
    while r <= last_row:
        no = ws.cell(r, 1).value
        lokasi = ws.cell(r, 2).value
        no_sp = ws.cell(r, 3).value
        tgl_sp = ws.cell(r, 4).value
        termin = ws.cell(r, 5).value
        bulan = ws.cell(r, 6).value
        nominal = ws.cell(r, 7).value
        ket = ws.cell(r, 8).value
        status = ws.cell(r, 9).value

        is_note = isinstance(no_sp, str) and re.match(r"^\d+\.", no_sp) and not lokasi and not bulan
        if is_note:
            notes.append(no_sp)
            blank_streak = 0
            r += 1
            continue

        is_blank = not no and not lokasi and not no_sp and not bulan and not nominal
        if is_blank:
            blank_streak += 1
            r += 1
            if blank_streak > 30:
                break
            continue
        blank_streak = 0

        if lokasi:
            current = {
                "no": no,
                "lokasi": str(lokasi).strip(),
                "noSp": str(no_sp).strip() if no_sp else None,
                "tanggalSp": iso(tgl_sp),
                "jumlahTermin": termin,
                "perBulan": [],
            }
            tagihan.append(current)
            if bulan:
                current["perBulan"].append({
                    "bulan": bulan, "nominal": nominal,
                    "keterangan": str(ket).strip() if ket else None,
                    "status": str(status).strip() if status else None,
                })
        elif bulan and current is not None:
            current["perBulan"].append({
                "bulan": bulan, "nominal": nominal,
                "keterangan": str(ket).strip() if ket else None,
                "status": str(status).strip() if status else None,
            })
        r += 1

    for t in tagihan:
        total = sum(b["nominal"] for b in t["perBulan"] if b["nominal"])
        selesai = sum(1 for b in t["perBulan"] if b["status"] and b["status"].upper() == "SELESAI")
        t["totalNominal"] = total
        t["totalBulan"] = len(t["perBulan"])
        t["selesai"] = selesai
        if t["totalBulan"] == 0:
            t["status"] = "belum"
        elif selesai == t["totalBulan"]:
            t["status"] = "selesai"
        elif selesai > 0:
            t["status"] = "proses"
        else:
            t["status"] = "belum"

    return tagihan, notes


def find_col(header, name, start=0):
    """Cari index (0-based) kolom header yang cocok persis (case-insensitive)
    dengan `name`, mulai dari index `start`. Dipakai supaya posisi kolom
    tidak perlu di-hardcode — tetap kebaca walau user menyisipkan/menggeser
    kolom. `start` dipakai untuk header yang muncul dua kali (mis. "BAPP"
    sebagai checklist boolean di awal, lalu "BAPP" lagi sebagai nomor berita
    acara belakangan) supaya ambil kemunculan yang benar."""
    name_up = name.strip().upper()
    for i in range(start, len(header)):
        v = header[i]
        if isinstance(v, str) and v.strip().upper() == name_up:
            return i
    return None


def split_lines(v):
    if v is None or v == "":
        return []
    if isinstance(v, (int, float)):
        return [v]
    return [s.strip() for s in str(v).split("\n") if s.strip()]


def parse_rupiah(s):
    digits = re.sub(r"[^\d]", "", s)
    return int(digits) if digits else None


def parse_bagian_nominal(bagian_val, nominal_val):
    """BAGIAN KERJA & NOMINAL TAGIHAN bisa berisi beberapa baris (dipisah
    newline) kalau satu lokasi punya beberapa bagian kerja sekaligus — tiap
    baris dipasangkan berurutan (baris ke-1 BAGIAN KERJA <-> baris ke-1
    NOMINAL TAGIHAN, dst). NOMINAL TAGIHAN bisa berupa angka polos atau teks
    "Rp. 6.577.423"."""
    bagian_list = split_lines(bagian_val)
    nominal_list = []
    for n in split_lines(nominal_val):
        nominal_list.append(int(round(n)) if isinstance(n, (int, float)) else parse_rupiah(str(n)))
    pairs = []
    for i in range(max(len(bagian_list), len(nominal_list))):
        pairs.append({
            "bagianKerja": str(bagian_list[i]) if i < len(bagian_list) else None,
            "nominal": nominal_list[i] if i < len(nominal_list) else None,
        })
    return pairs


def cell_at(ws, r, idx):
    """idx 0-based (hasil find_col) -> baca sel di baris r, atau None kalau
    kolomnya tidak ditemukan di sheet ini (mis. bulan Juni belum punya blok
    kontrak/SP)."""
    return ws.cell(r, idx + 1).value if idx is not None else None


def parse_dokumen_sheet(ws, code, label):
    header = [ws.cell(4, c).value for c in range(1, 55)]
    tgl_kirim_idx = find_col(header, "TANGGAL KIRIM")
    bagian_idx = find_col(header, "BAGIAN KERJA")
    nominal_idx = find_col(header, "NOMINAL TAGIHAN")
    jenis_dokumen = [h for h in header[3:tgl_kirim_idx] if h] if tgl_kirim_idx is not None else [h for h in header[3:18] if h]

    # Kolom detail kontrak/berita acara (baru dipakai mulai Agustus 2026) —
    # dicari mulai SETELAH TANGGAL KIRIM, karena "BAPP"/"BAST"/"BAP" juga
    # jadi nama kolom checklist boolean di awal sheet.
    search_from = (tgl_kirim_idx + 1) if tgl_kirim_idx is not None else 0
    bapp_idx = find_col(header, "BAPP", search_from)
    tgl_bapp_idx = find_col(header, "TANGGAL BAPP", search_from)
    bast_idx = find_col(header, "BAST", search_from)
    tgl_bast_idx = find_col(header, "TANGGAL BAST", search_from)
    bap_idx = find_col(header, "BAP", search_from)
    tgl_bap_idx = find_col(header, "TANGGAL BAP", search_from)
    no_inv_idx = find_col(header, "NO INVOICE DAN KWITANSI", search_from)
    no_sp_idx = find_col(header, "No SP", search_from)
    tgl_sp_idx = find_col(header, "Tanggal SP", search_from)
    lama_kontrak_idx = find_col(header, "LAMA KONTRAK", search_from)
    periode_kontrak_idx = find_col(header, "PERIODE KONTRAK", search_from)
    termin_idx = find_col(header, "TERMIN", search_from)
    metode_idx = find_col(header, "METODE", search_from)
    nomor_idx = find_col(header, "NOMOR", search_from)
    link_nomor_idx = find_col(header, "LINK NOMOR", search_from)

    lokasi_list = []
    r = 5
    last_row = ws.max_row
    while r <= last_row:
        lokasi = ws.cell(r, 2).value
        if not lokasi:
            break
        dokumen = {}
        lengkap = 0
        for i, jenis in enumerate(jenis_dokumen):
            val = ws.cell(r, 4 + i).value is True
            dokumen[jenis] = val
            if val:
                lengkap += 1
        tagihan = parse_bagian_nominal(
            ws.cell(r, bagian_idx + 1).value if bagian_idx is not None else None,
            ws.cell(r, nominal_idx + 1).value if nominal_idx is not None else None,
        )
        lokasi_list.append({
            "lokasi": str(lokasi).strip(),
            "dokumen": dokumen,
            "tanggalKirim": iso(ws.cell(r, tgl_kirim_idx + 1).value) if tgl_kirim_idx is not None else None,
            "totalDokumen": len(jenis_dokumen),
            "dokumenLengkap": lengkap,
            "pctLengkap": round(lengkap / len(jenis_dokumen), 4) if jenis_dokumen else None,
            "tagihan": tagihan,
            "totalNominal": sum(p["nominal"] for p in tagihan if p["nominal"]),
            "noInvoiceKwitansi": cell_display(cell_at(ws, r, no_inv_idx)),
            "bapp": {"nomor": cell_display(cell_at(ws, r, bapp_idx)), "tanggal": cell_display(cell_at(ws, r, tgl_bapp_idx))},
            "bast": {"nomor": cell_display(cell_at(ws, r, bast_idx)), "tanggal": cell_display(cell_at(ws, r, tgl_bast_idx))},
            "bap": {"nomor": cell_display(cell_at(ws, r, bap_idx)), "tanggal": cell_display(cell_at(ws, r, tgl_bap_idx))},
            "kontrak": {
                "noSp": cell_display(cell_at(ws, r, no_sp_idx)),
                "tanggalSp": cell_display(cell_at(ws, r, tgl_sp_idx)),
                "lamaKontrak": cell_display(cell_at(ws, r, lama_kontrak_idx)),
                "periodeKontrak": cell_display(cell_at(ws, r, periode_kontrak_idx)),
                "termin": cell_display(cell_at(ws, r, termin_idx)),
                "metode": cell_display(cell_at(ws, r, metode_idx)),
                "statusNomor": cell_display(cell_at(ws, r, nomor_idx)),
                "linkNomor": cell_display(cell_at(ws, r, link_nomor_idx)),
            },
        })
        r += 1

    pct_list = [l["pctLengkap"] for l in lokasi_list if l["pctLengkap"] is not None]
    return {
        "code": code,
        "label": label,
        "jenisDokumen": jenis_dokumen,
        "lokasi": lokasi_list,
        "pctRataRata": round(sum(pct_list) / len(pct_list), 4) if pct_list else None,
        "totalNominal": sum(l["totalNominal"] for l in lokasi_list),
    }


def build_invoice_dataset(invoice_src):
    empty = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sourceFile": None,
        "tersedia": False,
        "tagihan": [],
        "totalLokasiTagihan": 0,
        "totalNominalKeseluruhan": 0,
        "totalNominalBelumSelesai": 0,
        "totalNominalTagihan": 0,
        "dokumenBulanan": [],
        "catatan": [],
    }
    if not invoice_src:
        return empty

    wb = openpyxl.load_workbook(invoice_src, data_only=True)
    tagihan, notes = parse_tagihan_sheet(wb["MONITORING TAGIHAN"] if "MONITORING TAGIHAN" in wb.sheetnames else None)

    dokumen_bulanan = []
    for sn in wb.sheetnames:
        m = re.match(r"^([A-Za-z]+)\s+(\d{4})$", sn.strip())
        if not m:
            continue
        month_num = MONTH_NUM_ID.get(m.group(1).upper())
        if not month_num:
            continue
        code = f"{m.group(2)}-{month_num:02d}"
        dokumen_bulanan.append(parse_dokumen_sheet(wb[sn], code, f"{MONTH_LABEL_ID[month_num]} {m.group(2)}"))
    dokumen_bulanan.sort(key=lambda d: d["code"])

    total_nominal = sum(t["totalNominal"] for t in tagihan)
    total_belum = sum(
        b["nominal"] for t in tagihan for b in t["perBulan"]
        if b["nominal"] and (not b["status"] or b["status"].upper() != "SELESAI")
    )

    return {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sourceFile": invoice_src.split("/")[-1],
        "tersedia": True,
        "tagihan": tagihan,
        "totalLokasiTagihan": len(tagihan),
        "totalNominalKeseluruhan": total_nominal,
        "totalNominalBelumSelesai": total_belum,
        "totalNominalTagihan": sum(m["totalNominal"] for m in dokumen_bulanan),
        "dokumenBulanan": dokumen_bulanan,
        "catatan": notes,
    }


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "JURNAL HARIAN DINI SAFFANAH 2026.xlsx"
    invoice_src = sys.argv[2] if len(sys.argv) > 2 else None
    wb = openpyxl.load_workbook(src, data_only=True)
    foto_count_map = build_foto_count_map(wb)

    all_tasks = []
    for sn in wb.sheetnames:
        base = sn.strip().upper()
        if base.endswith(" FOTO") or base not in MONTH_NAMES_UPPER:
            continue
        all_tasks.extend(parse_month_sheet(wb[sn], foto_count_map))

    real_tasks = [t for t in all_tasks if t["tugas"] and str(t["tugas"]).strip().upper() != PLACEHOLDER_TASK]
    days = build_days(all_tasks)
    months = build_months(days)

    total_tugas = len(real_tasks)
    total_selesai = sum(1 for t in real_tasks if t["selesai"])
    jam_list = [d["jamKerjaJam"] for d in days if d["jamKerjaJam"] is not None]
    tanggal_list = [d["tanggal"] for d in days]

    invoice = build_invoice_dataset(invoice_src)

    dataset = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sourceFile": src.split("/")[-1],
        "months": months,
        "days": days,
        "tasks": real_tasks,
        "totalHariTercatat": len(days),
        "totalTugas": total_tugas,
        "totalSelesai": total_selesai,
        "pctSelesaiKeseluruhan": round(total_selesai / total_tugas, 4) if total_tugas else None,
        "totalFoto": sum(d["jumlahFoto"] for d in days),
        "rataJamKerjaKeseluruhan": round(sum(jam_list) / len(jam_list), 2) if jam_list else None,
        "rentangTanggal": {
            "mulai": min(tanggal_list) if tanggal_list else None,
            "akhir": max(tanggal_list) if tanggal_list else None,
        },
        "invoice": invoice,
    }

    out_path = "src/data/jurnal-data.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(dataset, f, ensure_ascii=False, indent=2)
    print(f"wrote {out_path}")
    print(f"  bulan: {len(months)}")
    print(f"  hari tercatat: {len(days)}")
    print(f"  total tugas: {total_tugas} (selesai {total_selesai})")
    if invoice["tersedia"]:
        print(f"  invoice: {invoice['totalLokasiTagihan']} lokasi tagihan, {len(invoice['dokumenBulanan'])} bulan dokumen")


if __name__ == "__main__":
    main()
